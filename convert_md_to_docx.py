import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    parse_inline_formatting(p, text)
    return p

def parse_inline_formatting(paragraph, text):
    """Parses bold, italic, inline code, and links in markdown text."""
    # Pattern to match bold, italic, code, links
    pattern = re.compile(
        r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    )
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # Ruby/Red code style
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                link_text, link_url = m.groups()
                run = paragraph.add_run(f"{link_text} ({link_url})")
                run.font.color.rgb = RGBColor(26, 115, 232)
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def convert_md_file_to_docx(md_path, docx_path):
    print(f"Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")
    with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()

    doc = Document()
    
    # Configure page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styles configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)

    i = 0
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle Code Blocks (``` ... ```)
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_block_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                # Render code block as a single-cell shaded table
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F6F8")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                # Header if language is present
                if code_block_lang:
                    hdr_p = cell.paragraphs[0]
                    hdr_p.paragraph_format.space_after = Pt(4)
                    run_hdr = hdr_p.add_run(f"[{code_block_lang.upper()}]")
                    run_hdr.bold = True
                    run_hdr.font.size = Pt(8.5)
                    run_hdr.font.color.rgb = RGBColor(108, 117, 125)
                    start_p_idx = 1
                else:
                    start_p_idx = 0

                content_p = cell.add_paragraph() if start_p_idx == 1 else cell.paragraphs[0]
                content_p.paragraph_format.space_before = Pt(0)
                content_p.paragraph_format.space_after = Pt(0)
                content_p.paragraph_format.line_spacing = Pt(13)
                
                raw_code = "\n".join(code_block_lines)
                code_run = content_p.add_run(raw_code)
                code_run.font.name = 'Consolas'
                code_run.font.size = Pt(9.5)
                code_run.font.color.rgb = RGBColor(40, 44, 52)

                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                i += 1
                continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Tables (| col1 | col2 |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            # Check if table continues
            if i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                continue
            else:
                # End of table, let's process table_lines
                parse_and_add_table(doc, table_lines)
                table_lines = []
                continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal Rule (---, ***, ___)
        if re.match(r'^(---|___|\*\*\*)$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D1D5DB"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        h_match = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if h_match:
            level = len(h_match.group(1))
            heading_text = h_match.group(2).strip()

            if level == 1:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.bold = True
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor(17, 24, 39) # Deep Navy / Slate 900
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = RGBColor(30, 58, 138) # Royal Blue / Indigo
            elif level == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.bold = True
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(55, 65, 81) # Slate 700
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.bold = True
                run.font.size = Pt(11.5)
                run.font.color.rgb = RGBColor(75, 85, 99)
            i += 1
            continue

        # Blockquote (> ...)
        if stripped.startswith('>'):
            quote_text = stripped.lstrip('> ').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("“ " + quote_text + " ”")
            run.italic = True
            run.font.color.rgb = RGBColor(75, 85, 99)
            i += 1
            continue

        # Unordered Bullet List (- or *)
        bullet_match = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_text = bullet_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if indent_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.5)
            parse_inline_formatting(p, bullet_text)
            i += 1
            continue

        # Numbered List (1. 2. etc.)
        num_match = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if num_match:
            num_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, num_text)
            i += 1
            continue

        # Standard Paragraph
        add_styled_paragraph(doc, stripped, space_after=6)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")

def parse_and_add_table(doc, raw_lines):
    """Converts markdown table lines to a styled Word table."""
    rows_data = []
    for line in raw_lines:
        cells = [c.strip() for c in line.strip('|').split('|')]
        # Skip divider row (--- | ---)
        if all(re.match(r'^\s*:?-+:?\s*$', c) for c in cells if c):
            continue
        rows_data.append(cells)

    if not rows_data:
        return

    num_cols = max(len(r) for r in rows_data)
    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True

    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = tbl.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(13)
            
            if is_header:
                set_cell_background(cell, "2563EB") # Indigo / Blue Accent
                set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                run = p.add_run(val)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255) # White text
                run.font.size = Pt(10)
            else:
                # Zebra striping
                bg_color = "F9FAFB" if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                parse_inline_formatting(p, val)
                for run in p.runs:
                    run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def main():
    workspace_root = os.path.dirname(os.path.abspath(__file__))
    docs_dir = os.path.join(workspace_root, "docs")
    output_dir = os.path.join(docs_dir, "word_documents")
    os.makedirs(output_dir, exist_ok=True)

    print("=" * 70)
    print("   RENTORA: BATCH MARKDOWN TO WORD (.DOCX) CONVERTER")
    print("=" * 70)

    count = 0
    for root, dirs, files in os.walk(docs_dir):
        # Skip output dir to avoid loops
        if "word_documents" in root:
            continue
        for file in files:
            if file.endswith('.md'):
                md_full_path = os.path.join(root, file)
                
                # Relative path from docs to mirror structure
                rel_path = os.path.relpath(md_full_path, docs_dir)
                rel_base, _ = os.path.splitext(rel_path)
                
                # 1. Save in mirroring folder inside docs/word_documents/
                target_docx_central = os.path.join(output_dir, rel_base + ".docx")
                os.makedirs(os.path.dirname(target_docx_central), exist_ok=True)
                
                # 2. Also save directly next to the original .md file for convenience
                target_docx_local = os.path.splitext(md_full_path)[0] + ".docx"

                try:
                    convert_md_file_to_docx(md_full_path, target_docx_central)
                    # Copy or save locally as well
                    import shutil
                    shutil.copy2(target_docx_central, target_docx_local)
                    count += 1
                except Exception as e:
                    print(f"Error converting {file}: {e}")

    print("=" * 70)
    print(f"Successfully converted {count} Markdown files to formatted Word documents!")
    print(f"Centralized Word Docs Folder: {output_dir}")
    print("=" * 70)

if __name__ == '__main__':
    main()
